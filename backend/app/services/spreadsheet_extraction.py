"""Bounded Excel-to-text extraction for task search; never evaluates formulas."""

from __future__ import annotations

from contextlib import ExitStack
from datetime import date, datetime, time
import logging
import os
from zipfile import ZipFile

import openpyxl
from openpyxl.utils import get_column_letter
import xlrd

logger = logging.getLogger(__name__)

SPREADSHEET_EXTENSIONS = (".xls", ".xlsx")
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_SHEETS = 20
MAX_ROWS = 10_000  # Across the workbook, including empty rows.
MAX_COLUMNS = 256
MAX_CELLS = 200_000
MAX_CELL_CHARS = 400
MAX_DOCUMENTS = 400
MAX_DOCUMENT_CHARS = 1000  # Fits the sync pipeline's embedding chunk size.


class ExtractionLimit(ValueError):
    pass


def is_spreadsheet(filename: str) -> bool:
    return filename.lower().endswith(SPREADSHEET_EXTENSIONS)


def _format_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float):
        return format(value, ".15g")
    return " ".join(str(value).replace("\x00", "").split())


def _xlsx_rows(path: str, warnings: set[str]):
    with ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise ExtractionLimit("Workbook exceeds the uncompressed size limit.")

    # Read both cached results and formulas, so formulas without a cached value
    # are reported explicitly rather than silently disappearing from the evidence.
    with ExitStack() as stack:
        books = []
        for data_only in (True, False):
            source = stack.enter_context(open(path, "rb"))  # Temp paths have no extension.
            book = openpyxl.load_workbook(
                source, read_only=True, data_only=data_only, keep_links=False,
            )
            stack.callback(book.close)
            books.append(book)
        values_book, formulas_book = books
        for sheet_index, values_sheet in enumerate(values_book.worksheets):
            if sheet_index >= MAX_SHEETS:
                raise ExtractionLimit("Workbook exceeds the worksheet limit.")
            formulas_sheet = formulas_book[values_sheet.title]
            # Some producers incorrectly declare A1:A1 as the used range.
            values_sheet.reset_dimensions()
            formulas_sheet.reset_dimensions()
            for row_number, (values, formulas) in enumerate(
                zip(values_sheet.iter_rows(), formulas_sheet.iter_rows()), start=1,
            ):
                if len(values) > MAX_COLUMNS:
                    warnings.add("Columns beyond the column limit were omitted.")
                cells = []
                for value_cell, formula_cell in zip(values[:MAX_COLUMNS], formulas[:MAX_COLUMNS]):
                    value = _format_value(value_cell.value)
                    if formula_cell.data_type == "f" and value_cell.value is None:
                        formula = _format_value(formula_cell.value)
                        value = f"Formula {formula} [cached result unavailable; not calculated]"
                    cells.append(value)
                yield values_sheet.title, row_number, cells


def _xls_rows(path: str, warnings: set[str]):
    book = xlrd.open_workbook(path, on_demand=True, ragged_rows=True)
    try:
        for sheet_index in range(book.nsheets):
            if sheet_index >= MAX_SHEETS:
                raise ExtractionLimit("Workbook exceeds the worksheet limit.")
            sheet = book.sheet_by_index(sheet_index)
            try:
                if sheet.ncols > MAX_COLUMNS:
                    warnings.add("Columns beyond the column limit were omitted.")
                for row_index in range(sheet.nrows):
                    cells = []
                    for col_index in range(min(sheet.row_len(row_index), MAX_COLUMNS)):
                        cell = sheet.cell(row_index, col_index)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode)
                            if 0 <= cell.value < 1:
                                value = value.time()
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code.get(value, "#ERROR")
                        cells.append(_format_value(value))
                    yield sheet.name, row_index + 1, cells
            finally:
                book.unload_sheet(sheet_index)
    finally:
        book.release_resources()


def extract_spreadsheet_documents(path: str, filename: str) -> list[dict[str, str]]:
    """Return bounded, source-labelled text chunks, including any extraction limits.

    Bad/encrypted workbooks remain downloadable: parsing failures become a source
    notice. Storage and embedding failures are handled by the caller, not hidden.
    XLS uses the results cached by Excel; XLSX also reports uncached formulas.
    """
    documents: list[dict[str, str]] = []
    warnings: set[str] = set()
    filename_label = _format_value(filename)[:180]
    current_sheet = None
    prefix = ""
    body = ""
    first_row = last_row = 0
    first_populated_row = ""

    def flush():
        nonlocal body
        if body:
            if len(documents) >= MAX_DOCUMENTS - 1:  # Reserve space for a notice.
                raise ExtractionLimit("Workbook exceeds the searchable text limit.")
            documents.append({
                "text": prefix + body,
                "section": f"sheet:{current_sheet}:rows:{first_row}-{last_row}",
            })
            body = ""

    try:
        if not is_spreadsheet(filename):
            raise ValueError("Unsupported spreadsheet extension")
        if os.path.getsize(path) > MAX_FILE_BYTES:
            raise ExtractionLimit("Workbook exceeds the file size limit.")
        rows = _xlsx_rows(path, warnings) if filename.lower().endswith(".xlsx") else _xls_rows(path, warnings)
        scanned_cells = 0
        with ExitStack() as stack:
            stack.callback(rows.close)
            for scanned_rows, (sheet, row_number, cells) in enumerate(rows, start=1):
                scanned_cells += len(cells)
                if scanned_rows > MAX_ROWS or scanned_cells > MAX_CELLS:
                    raise ExtractionLimit("Workbook exceeds the row/cell scan limit.")
                if sheet != current_sheet:
                    flush()
                    current_sheet = sheet
                    first_populated_row = ""
                entries = []
                for column, value in enumerate(cells, start=1):
                    if not value:
                        continue
                    if len(value) > MAX_CELL_CHARS:
                        warnings.add("Long cell values were truncated.")
                        value = value[:MAX_CELL_CHARS] + " [truncated]"
                    entries.append(f"{get_column_letter(column)}{row_number}={value}")
                if not entries:
                    continue
                if not first_populated_row:
                    # Repeat initial labels so later chunks retain table context.
                    first_populated_row = " | ".join(entries)[:200]
                prefix = (
                    f"Workbook: {filename_label}\nWorksheet: {_format_value(sheet)[:31]}\n"
                    f"First populated row (excerpt): {first_populated_row}\n"
                )
                for entry in entries:
                    separator = " | " if last_row == row_number else "\n"
                    if body and len(prefix + body + separator + entry) > MAX_DOCUMENT_CHARS:
                        flush()
                    if not body:
                        first_row = row_number
                    body += (separator if body else "") + entry
                    last_row = row_number
        flush()
    except ExtractionLimit as exc:
        warnings.add(str(exc))
        if body and len(documents) < MAX_DOCUMENTS - 1:
            flush()
    except Exception:
        logger.exception("[SPREADSHEET] Could not read %s", filename)
        # Keep already completed chunks but do not pass partially parsed rows on.
        warnings.add("Workbook could not be fully read; it may be corrupt, encrypted, or unsupported.")

    if warnings or not documents:
        detail = " ".join(sorted(warnings)) if warnings else "No populated worksheet cells were found."
        documents.append({
            "text": (
                f"Workbook: {filename_label}\nExtraction notice: {detail} "
                "Extracted evidence may be incomplete. Consult the original workbook for full details."
            ),
            "section": "spreadsheet:extraction-notice",
        })
        logger.warning("[SPREADSHEET] %s: %s", filename, detail)
    return documents
