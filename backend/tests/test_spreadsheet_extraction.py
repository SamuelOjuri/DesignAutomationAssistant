from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

import openpyxl
import pytest

from backend.app.services import spreadsheet_extraction as excel


XLS_FIXTURE = Path(__file__).parent / "fixtures" / "spreadsheets" / "pricing_schedule.xls"


def make_xlsx(path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Pricing"
    for row in [
        ["Item", "Quantity", "Unit price"],
        ["Roof insulation", 12, 25.5],
        ["Zero quantity", 0, False],
        ["Received", datetime(2026, 9, 11)],
    ]:
        sheet.append(row)
    book.create_sheet("Notes")["A3"] = "Delivery included"
    book.save(path)
    book.close()


@pytest.mark.parametrize("extension", [".xls", ".xlsx"])
def test_reads_real_workbooks_with_cell_and_sheet_references(tmp_path, extension):
    path = tmp_path / "download-without-extension"
    if extension == ".xls":
        path.write_bytes(XLS_FIXTURE.read_bytes())
    else:
        make_xlsx(path)
    docs = excel.extract_spreadsheet_documents(str(path), "Pricing Schedule" + extension.upper())
    pricing = "\n".join(d["text"] for d in docs if "Worksheet: Pricing" in d["text"])
    assert "A2=Roof insulation" in pricing
    assert "B2=12" in pricing
    assert "C2=25.5" in pricing
    assert "B3=0" in pricing
    assert "C3=FALSE" in pricing
    assert "B4=2026-09-11T00:00:00" in pricing
    assert any(d["section"] == "sheet:Notes:rows:3-3" and "A3=Delivery included" in d["text"] for d in docs)
    assert all("Workbook: Pricing Schedule" in d["text"] for d in docs)
    assert all(len(d["text"]) <= excel.MAX_DOCUMENT_CHARS for d in docs)
    # In particular, read-only workbook handles must be closed on Windows.
    path.unlink()


def test_cached_and_uncached_formulas(tmp_path):
    path = tmp_path / "formulas.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["Cached total", "=12*25.5", "=1+1"])
    book.save(path)
    book.close()
    # openpyxl does not calculate formulas; emulate Excel's saved cached result.
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = ElementTree.fromstring(entries["xl/worksheets/sheet1.xml"])
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root.find(".//s:c[@r='B1']/s:v", ns).text = "306"
    entries["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root)
    with ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    text = "\n".join(d["text"] for d in excel.extract_spreadsheet_documents(str(path), path.name))
    assert "B1=306" in text
    assert "C1=Formula =1+1 [cached result unavailable; not calculated]" in text


def test_chunks_repeat_workbook_and_initial_labels_without_losing_rows(tmp_path):
    path = tmp_path / "many-rows.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Prices"
    book.active.append(["Product", "Price"])
    for number in range(2, 202):
        book.active.append([f"Product {number}", number * 10])
    book.save(path)
    book.close()
    docs = excel.extract_spreadsheet_documents(str(path), path.name)
    assert len(docs) > 1
    for doc in docs:
        assert "Workbook: many-rows.xlsx" in doc["text"]
        assert "Worksheet: Prices" in doc["text"]
        assert "A1=Product | B1=Price" in doc["text"]
        assert len(doc["text"]) <= excel.MAX_DOCUMENT_CHARS
    assert "B201=2010" in docs[-1]["text"]


@pytest.mark.parametrize("limit,value,notice", [
    ("MAX_FILE_BYTES", 1, "file size limit"),
    ("MAX_UNCOMPRESSED_BYTES", 1, "uncompressed size limit"),
    ("MAX_ROWS", 2, "row/cell scan limit"),
    ("MAX_CELLS", 3, "row/cell scan limit"),
    ("MAX_SHEETS", 1, "worksheet limit"),
    ("MAX_COLUMNS", 2, "column limit"),
    ("MAX_CELL_CHARS", 5, "truncated"),
])
def test_limits_are_explicit_and_close_workbook(tmp_path, monkeypatch, limit, value, notice):
    path = tmp_path / "limits.xlsx"
    make_xlsx(path)
    monkeypatch.setattr(excel, limit, value)
    docs = excel.extract_spreadsheet_documents(str(path), path.name)
    assert notice in docs[-1]["text"]
    assert docs[-1]["section"] == "spreadsheet:extraction-notice"
    path.unlink()


def test_document_limit_preserves_prior_chunks_and_reports_partial_extraction(tmp_path, monkeypatch):
    path = tmp_path / "large.xlsx"
    book = openpyxl.Workbook()
    for _ in range(20):
        book.active.append(["x" * 300] * 5)
    book.save(path)
    book.close()
    monkeypatch.setattr(excel, "MAX_DOCUMENTS", 3)
    docs = excel.extract_spreadsheet_documents(str(path), path.name)
    assert len(docs) == 3
    assert "searchable text limit" in docs[-1]["text"]
    assert all(len(d["text"]) <= excel.MAX_DOCUMENT_CHARS for d in docs)
    path.unlink()


@pytest.mark.parametrize("extension", [".xls", ".xlsx"])
def test_corrupt_workbook_returns_notice(tmp_path, extension):
    path = tmp_path / ("broken" + extension)
    path.write_bytes(b"not an Excel workbook")
    docs = excel.extract_spreadsheet_documents(str(path), path.name)
    assert len(docs) == 1
    assert "could not be fully read" in docs[0]["text"]


def test_empty_workbook_is_reported(tmp_path):
    path = tmp_path / "empty.xlsx"
    book = openpyxl.Workbook()
    book.save(path)
    book.close()
    docs = excel.extract_spreadsheet_documents(str(path), path.name)
    assert "No populated worksheet cells" in docs[0]["text"]


def test_incorrect_xlsx_dimensions_do_not_hide_cells(tmp_path):
    path = tmp_path / "dimensions.xlsx"
    make_xlsx(path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = ElementTree.fromstring(entries["xl/worksheets/sheet1.xml"])
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root.find("s:dimension", ns).set("ref", "A1:A1")
    entries["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root)
    with ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    docs = excel.extract_spreadsheet_documents(str(path), path.name)
    assert any("C2=25.5" in doc["text"] for doc in docs)


def test_xls_limit_closes_mapped_file(tmp_path, monkeypatch):
    path = tmp_path / "limit.xls"
    path.write_bytes(XLS_FIXTURE.read_bytes())
    monkeypatch.setattr(excel, "MAX_ROWS", 2)
    docs = excel.extract_spreadsheet_documents(str(path), path.name)
    assert "row/cell scan limit" in docs[-1]["text"]
    assert any("Roof insulation" in doc["text"] for doc in docs[:-1])
    path.unlink()
